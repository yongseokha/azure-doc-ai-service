"""검증 원문 기반 약관 검증(/terms/verify-ask)용 리포트.

시트 구성/스타일은 기존 약관 검증 리포트(report_service)와 같고, 항목 컬럼만 다르다.
value 분해(subClaim)가 없어서 항목 상세 표에 세로 셀병합이 필요 없다.
"""

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.schemas.terms import TermsAskVerificationResult
from app.services.report_service import (
    STATUS_LABELS,
    SUMMARY_COLUMN_WIDTHS,
    SUMMARY_MAX_COL,
    _stats_from_items,
    _write_detailed_stats,
    _write_result_summary,
    _write_section_title,
    _write_table,
    _write_title,
    _write_verification_info,
)

ITEM_COLUMNS = ["itemNm", "askValue", "llmValue", "evidence", "page", "article", "reason", "matchRate", "status"]
COLUMN_WIDTHS = [22, 24, 24, 38, 12, 16, 32, 12, 14]  # ITEM_COLUMNS 순서
MAX_COL = len(COLUMN_WIDTHS)


def _flatten(result: TermsAskVerificationResult) -> pd.DataFrame:
    rows = [
        {
            "name": name_result.name,
            "termNm": doc_result.termNm,
            "nameMatchRate": doc_result.nameMatchRate,
            **{col: getattr(item, col) for col in ITEM_COLUMNS},
        }
        for name_result in result.data
        for doc_result in name_result.documents
        for item in doc_result.items
    ]
    return pd.DataFrame(rows, columns=["name", "termNm", "nameMatchRate", *ITEM_COLUMNS])


def _build_item_row(item) -> dict:
    row = {col: getattr(item, col) for col in ITEM_COLUMNS}
    row["status"] = STATUS_LABELS.get(row["status"], row["status"])
    return row


def compute_overview_stats(result: TermsAskVerificationResult) -> dict:
    all_items = [item for name_result in result.data for doc_result in name_result.documents for item in doc_result.items]
    return _stats_from_items(all_items)


def build_result_payload(result: TermsAskVerificationResult, verified_at: str) -> dict:
    """콜백의 vrfDataResltJson과 GET 상태조회가 공유하는 전체 결과 payload.
    기존 약관 검증과 같은 구조(전체/name별/문서별 통계 인라인)이고 items 필드만 다르다."""
    data = []
    for name_result in result.data:
        name_items = [item for doc_result in name_result.documents for item in doc_result.items]
        documents = [
            {
                "ocrResltKey": doc_result.ocrResltKey,
                "termNm": doc_result.termNm,
                "aplyDate": doc_result.aplyDate,
                "nameMatchRate": doc_result.nameMatchRate,
                **_stats_from_items(doc_result.items),
                "items": [item.model_dump() for item in doc_result.items],
            }
            for doc_result in name_result.documents
        ]
        data.append({"name": name_result.name, **_stats_from_items(name_items), "documents": documents})

    return {
        "knwlgNm": result.knwlgNm,
        "verifiedAt": verified_at,
        **compute_overview_stats(result),
        "data": data,
    }


def _write_item_detail(ws, result: TermsAskVerificationResult, start_row: int, max_col: int) -> int:
    row = _write_section_title(ws, start_row, "항목별 상세 비교 결과", max_col)
    for name_result in result.data:
        for doc_result in name_result.documents:
            item_df = pd.DataFrame([_build_item_row(item) for item in doc_result.items], columns=ITEM_COLUMNS)
            row = _write_table(
                ws,
                item_df,
                row,
                max_col,
                f"{name_result.name} - {doc_result.termNm}",
                highlight_header=True,
                status_column="status",
                bold_columns=frozenset({"itemNm"}),
            )
    return row


def build_report_xlsx(result: TermsAskVerificationResult, verified_at: str | None = None) -> bytes:
    """기존 리포트와 같은 2개 시트("검증 요약" / "상세 결과") 구성의 xlsx를 만든다."""
    verified_at = verified_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "검증 요약"
    row = 1
    row = _write_title(ws1, row, SUMMARY_MAX_COL)
    row = _write_verification_info(ws1, row, SUMMARY_MAX_COL, verified_at, result)
    _write_result_summary(ws1, row, SUMMARY_MAX_COL, result)
    for col_idx, width in enumerate(SUMMARY_COLUMN_WIDTHS, start=1):
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    ws2 = wb.create_sheet("상세 결과")
    row = 1
    row = _write_detailed_stats(ws2, _flatten(result), row, MAX_COL)
    _write_item_detail(ws2, result, row, MAX_COL)
    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
